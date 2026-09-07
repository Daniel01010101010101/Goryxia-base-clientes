"""Prueba de punta a punta del pipeline usando la fixture sintetica.

No sale a la red: alimenta el pipeline con negocios ya parseados y verifica
que la deduplicacion, el scoring, el orden y los cuatro exportadores funcionan.
"""
import json
import tempfile
import unittest
from pathlib import Path

from goryxia.config import COLUMNAS, Ajustes
from goryxia.exporters import (
    exportar_csv,
    exportar_excel,
    exportar_geojson,
    exportar_pdf,
)
from goryxia.pipeline import ejecutar
from goryxia.sources.overpass import elemento_a_negocio

FIXTURE = Path(__file__).parent / "fixtures" / "overpass_demo.json"


def cargar_fixture():
    datos = json.loads(FIXTURE.read_text(encoding="utf-8"))
    negocios = [elemento_a_negocio(e) for e in datos["elements"]]
    return [n for n in negocios if n is not None]


class TestPipelineCompleto(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.crudos = cargar_fixture()
        ajustes = Ajustes(scrapear_webs=False, usar_datos_abiertos=False,
                          usar_google=False)
        cls.resultado = ejecutar(ajustes, negocios_precargados=cls.crudos,
                                 exportar=False)

    def test_la_fixture_produce_negocios(self):
        self.assertGreater(len(self.crudos), 300)

    def test_deduplica_los_casos_sembrados(self):
        total = sum(self.resultado.duplicados.values())
        self.assertGreaterEqual(total, 3)
        self.assertLess(len(self.resultado.negocios), len(self.crudos))

    def test_todos_los_registros_tienen_id_unico(self):
        ids = [n.id for n in self.resultado.negocios]
        self.assertTrue(all(ids))
        self.assertEqual(len(ids), len(set(ids)))

    def test_todos_tienen_score_prioridad_y_servicios(self):
        for negocio in self.resultado.negocios:
            self.assertIn(negocio.prioridad, {"ALTA", "MEDIA", "BAJA"})
            self.assertIn(negocio.presencia_digital, {"ALTA", "MEDIA", "BAJA"})
            self.assertIsInstance(negocio.score, int)
            self.assertTrue(negocio.servicios)
            self.assertTrue(negocio.observaciones)

    def test_la_base_queda_ordenada_por_contactabilidad(self):
        escalones = []
        for negocio in self.resultado.negocios:
            if negocio.tiene_celular:
                escalones.append(1)
            elif negocio.tiene_correo:
                escalones.append(3)
            elif negocio.tiene_web:
                escalones.append(4)
            else:
                escalones.append(5)
        self.assertEqual(escalones, sorted(escalones),
                         "Los negocios con celular deben ir primero")

    def test_los_que_tienen_celular_tienen_link_de_whatsapp(self):
        con_celular = [n for n in self.resultado.negocios if n.tiene_celular]
        self.assertGreater(len(con_celular), 0)
        for negocio in con_celular:
            self.assertTrue(negocio.link_whatsapp.startswith("https://wa.me/57"))

    def test_se_asigna_localidad_y_distancia(self):
        con_coordenadas = [n for n in self.resultado.negocios if n.lat is not None]
        self.assertTrue(all(n.localidad for n in con_coordenadas))
        self.assertTrue(all(n.distancia_km is not None for n in con_coordenadas))


class TestExportadores(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        crudos = cargar_fixture()
        ajustes = Ajustes(scrapear_webs=False, usar_datos_abiertos=False)
        cls.resultado = ejecutar(ajustes, negocios_precargados=crudos, exportar=False)
        cls.negocios = cls.resultado.negocios
        cls.temporal = tempfile.TemporaryDirectory()
        cls.dir = Path(cls.temporal.name)

    @classmethod
    def tearDownClass(cls):
        cls.temporal.cleanup()

    def test_csv_tiene_el_esquema_exacto(self):
        import csv

        ruta = exportar_csv(self.negocios, self.dir / "base.csv")
        with ruta.open(encoding="utf-8-sig") as archivo:
            lector = csv.DictReader(archivo)
            self.assertEqual(lector.fieldnames, COLUMNAS)
            filas = list(lector)
        self.assertEqual(len(filas), len(self.negocios))

    def test_excel_tiene_hojas_clientes_y_dashboard(self):
        from openpyxl import load_workbook

        ruta = exportar_excel(self.negocios, self.dir / "base.xlsx",
                              self.resultado.resumen)
        libro = load_workbook(ruta)
        self.assertIn("Clientes", libro.sheetnames)
        self.assertIn("Dashboard", libro.sheetnames)
        hoja = libro["Clientes"]
        self.assertEqual([c.value for c in hoja[1]], COLUMNAS)
        self.assertEqual(hoja.max_row, len(self.negocios) + 1)

    def test_geojson_es_valido(self):
        ruta = exportar_geojson(self.negocios, self.dir / "clientes.geojson")
        datos = json.loads(ruta.read_text(encoding="utf-8"))
        self.assertEqual(datos["type"], "FeatureCollection")
        self.assertGreater(len(datos["features"]), 0)
        primero = datos["features"][0]
        self.assertEqual(primero["geometry"]["type"], "Point")
        self.assertIn("Link WhatsApp", primero["properties"])

    def test_pdf_se_genera(self):
        ruta = exportar_pdf(self.negocios, self.dir / "base.pdf",
                            self.resultado.resumen, max_leads=80)
        self.assertTrue(ruta.exists())
        self.assertGreater(ruta.stat().st_size, 5000)
        self.assertEqual(ruta.read_bytes()[:4], b"%PDF")


if __name__ == "__main__":
    unittest.main()
